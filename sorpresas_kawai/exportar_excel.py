import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse


def aplicar_estilo_encabezado(celda, titulo, bg_color="D81B60", text_color="FFFFFF"):
    celda.value = titulo
    celda.font = Font(name="Segoe UI", size=11, bold=True, color=text_color)
    celda.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def aplicar_borde(celda, thin_border):
    celda.border = thin_border


def autoajustar_columnas(ws, min_width=12, max_width=50):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            lineas = val_str.split('\n')
            for l in lineas:
                if len(l) > max_len:
                    max_len = len(l)
        ancho = max(min(max_len + 4, max_width), min_width)
        ws.column_dimensions[col_letter].width = ancho


def generar_excel_metricas_ventas(
    pedidos_qs,
    resumen_data,
    productos_metricas,
    combos_metricas,
    alertas_stock,
    usuario_generador="Admin",
    rango_etiqueta="Todos los registros",
    estado_etiqueta="Todos los estados"
):
    """
    Genera un archivo Excel (.xlsx) profesional, visualmente atractivo y con múltiples hojas
    para el análisis gerencial de ventas e inventario de Sorpresas Kawaii.
    """
    wb = openpyxl.Workbook()
    
    # Paleta de colores Sorpresas Kawaii
    COLOR_PRIMARIO = "C2185B"     # Rosa Fucsia Profundo
    COLOR_SECUNDARIO = "7B1FA2"   # Púrpura Kawaii
    COLOR_TITULO = "880E4F"       # Borgoña Oscuro
    COLOR_FILA_PAR = "FCE4EC"     # Rosa Muy Suave
    COLOR_FILA_PAR_PURP = "F3E5F5" # Púrpura Muy Suave
    COLOR_ALERTA = "FFEBEE"       # Rojo Suave
    COLOR_ALERTA_TXT = "C62828"   # Rojo Fuerte
    
    thin_side = Side(border_style="thin", color="D1D5DB")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    # -------------------------------------------------------------
    # HOJA 1: RESUMEN GENERAL DE VENTAS
    # -------------------------------------------------------------
    ws_resumen = wb.active
    ws_resumen.title = "Resumen Ejecutivo"
    ws_resumen.views.sheetView[0].showGridLines = True
    
    # Título Principal
    ws_resumen.merge_cells("A1:F2")
    cell_title = ws_resumen["A1"]
    cell_title.value = "🌸 SORPRESAS KAWAII · REPORTE GERENCIAL DE VENTAS Y MÉTRICAS"
    cell_title.font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    cell_title.fill = PatternFill(start_color=COLOR_TITULO, end_color=COLOR_TITULO, fill_type="solid")
    cell_title.alignment = Alignment(horizontal="center", vertical="center")
    
    # Metadata del reporte
    ws_resumen["A4"] = "Fecha de Generación:"
    ws_resumen["B4"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws_resumen["A5"] = "Generado por:"
    ws_resumen["B5"] = str(usuario_generador)
    ws_resumen["A6"] = "Período Analizado:"
    ws_resumen["B6"] = rango_etiqueta
    ws_resumen["A7"] = "Filtro de Estado:"
    ws_resumen["B7"] = estado_etiqueta
    
    for row in range(4, 8):
        ws_resumen[f"A{row}"].font = Font(name="Segoe UI", bold=True, color="374151")
        ws_resumen[f"B{row}"].font = Font(name="Segoe UI", color="1F2937")
        ws_resumen[f"A{row}"].border = thin_border
        ws_resumen[f"B{row}"].border = thin_border

    # Encabezado sección KPIs
    ws_resumen.merge_cells("A9:D9")
    ws_resumen["A9"] = "INDICADORES CLAVE DE DESEMPEÑO (KPIs)"
    ws_resumen["A9"].font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
    ws_resumen["A9"].fill = PatternFill(start_color=COLOR_PRIMARIO, end_color=COLOR_PRIMARIO, fill_type="solid")
    ws_resumen["A9"].alignment = Alignment(horizontal="center", vertical="center")
    
    kpis = [
        ("Ingresos Totales Confirmados (Pagados/Enviados/Entregados)", resumen_data.get('ingresos_confirmados', 0.0), "$ #,##0"),
        ("Ingresos Brutos Potenciales (Incluye Pendientes)", resumen_data.get('ingresos_totales', 0.0), "$ #,##0"),
        ("Total de Órdenes / Pedidos Realizados", resumen_data.get('total_pedidos', 0), "#,##0"),
        ("Pedidos Confirmados y Pagados", resumen_data.get('pedidos_pagados', 0), "#,##0"),
        ("Pedidos en Estado Pendiente de Pago", resumen_data.get('pedidos_pendientes', 0), "#,##0"),
        ("Pedidos Cancelados", resumen_data.get('pedidos_cancelados', 0), "#,##0"),
        ("Ticket Promedio por Pedido Confirmado", resumen_data.get('ticket_promedio', 0.0), "$ #,##0"),
        ("Total Unidades de Productos Individuales Vendidas", resumen_data.get('unidades_productos_vendidas', 0), "#,##0"),
        ("Total Unidades de Combos Vendidas", resumen_data.get('unidades_combos_vendidas', 0), "#,##0"),
        ("Cantidad de Productos con Alerta de Stock Bajo (<= 5)", resumen_data.get('productos_alerta_stock', 0), "#,##0"),
    ]
    
    current_row = 10
    for metrica, valor, num_format in kpis:
        c1 = ws_resumen.cell(row=current_row, column=1, value=metrica)
        ws_resumen.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
        c2 = ws_resumen.cell(row=current_row, column=4, value=valor)
        
        c1.font = Font(name="Segoe UI", bold=True, color="1F2937")
        c2.font = Font(name="Segoe UI", bold=True, color=COLOR_PRIMARIO if "$" in num_format else "111827")
        c2.number_format = num_format
        c2.alignment = Alignment(horizontal="right", vertical="center")
        
        for col_i in range(1, 5):
            ws_resumen.cell(row=current_row, column=col_i).border = thin_border
            if current_row % 2 == 0:
                ws_resumen.cell(row=current_row, column=col_i).fill = PatternFill(start_color=COLOR_FILA_PAR, fill_type="solid")
        
        current_row += 1

    autoajustar_columnas(ws_resumen, min_width=15, max_width=60)
    
    # -------------------------------------------------------------
    # HOJA 2: DETALLE COMPLETO DE PEDIDOS
    # -------------------------------------------------------------
    ws_pedidos = wb.create_sheet(title="Detalle de Pedidos")
    ws_pedidos.views.sheetView[0].showGridLines = True
    
    headers_pedidos = [
        "N° Pedido", "ID Sistema", "Fecha y Hora", "Usuario Cliente",
        "Nombre Cliente", "Correo Electrónico", "Total Pagado ($ COP)",
        "Estado del Pedido", "Método de Pago", "Empresa de Envío", "Número de Guía"
    ]
    
    ws_pedidos.row_dimensions[1].height = 26
    for col_idx, h_text in enumerate(headers_pedidos, start=1):
        c = ws_pedidos.cell(row=1, column=col_idx)
        aplicar_estilo_encabezado(c, h_text, bg_color=COLOR_PRIMARIO)
        c.border = thin_border
        
    row_idx = 2
    for p in pedidos_qs:
        ws_pedidos.row_dimensions[row_idx].height = 20
        nombre_completo = f"{p.id_usuario.first_name} {p.id_usuario.last_name}".strip() or p.id_usuario.username
        
        datos_fila = [
            p.codigo_pedido,
            str(p.id),
            p.fecha_pedido.strftime("%Y-%m-%d %H:%M") if p.fecha_pedido else "",
            p.id_usuario.username,
            nombre_completo,
            p.id_usuario.email,
            p.valor_pagado,
            p.get_estado_display(),
            p.metodo_pago or "No especificado",
            p.empresa_envio or "—",
            p.numero_guia or "—"
        ]
        
        for col_idx, val in enumerate(datos_fila, start=1):
            cell = ws_pedidos.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            
            if col_idx == 7:
                cell.number_format = "$ #,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="047857")
            elif col_idx in (1, 3, 8):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
            if row_idx % 2 == 0:
                cell.fill = PatternFill(start_color=COLOR_FILA_PAR, fill_type="solid")
                
        row_idx += 1
        
    autoajustar_columnas(ws_pedidos, min_width=12, max_width=45)

    # -------------------------------------------------------------
    # HOJA 3: VENTAS POR PRODUCTO Y ROTACIÓN DE STOCK
    # -------------------------------------------------------------
    ws_prod = wb.create_sheet(title="Ventas por Producto")
    ws_prod.views.sheetView[0].showGridLines = True
    
    headers_prod = [
        "Producto", "Categoría", "Precio Unitario ($ COP)",
        "Unidades Vendidas", "Ingresos Generados ($ COP)",
        "Stock Actual", "Estado de Inventario"
    ]
    
    ws_prod.row_dimensions[1].height = 26
    for col_idx, h_text in enumerate(headers_prod, start=1):
        c = ws_prod.cell(row=1, column=col_idx)
        aplicar_estilo_encabezado(c, h_text, bg_color=COLOR_SECUNDARIO)
        c.border = thin_border

    row_idx = 2
    for item in productos_metricas:
        ws_prod.row_dimensions[row_idx].height = 20
        stock_val = item.get('stock', 0)
        if stock_val == 0:
            estado_stock = "⚠️ AGOTADO"
        elif stock_val <= 5:
            estado_stock = "⚠️ STOCK CRÍTICO"
        elif stock_val <= 10:
            estado_stock = "⚡ STOCK BAJO"
        else:
            estado_stock = "✅ ÓPTIMO"

        datos_fila = [
            item.get('nombre', ''),
            item.get('categoria', 'Sin Categoría'),
            item.get('valor_unitario', 0.0),
            item.get('unidades_vendidas', 0),
            item.get('ingresos_generados', 0.0),
            stock_val,
            estado_stock
        ]
        
        for col_idx, val in enumerate(datos_fila, start=1):
            cell = ws_prod.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            
            if col_idx in (3, 5):
                cell.number_format = "$ #,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (4, 6):
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.font = Font(name="Segoe UI", size=10, bold=True)
            elif col_idx == 7:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if "⚠️" in str(val):
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color=COLOR_ALERTA_TXT)
                    cell.fill = PatternFill(start_color=COLOR_ALERTA, fill_type="solid")
                else:
                    cell.font = Font(name="Segoe UI", size=10, color="065F46")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            if row_idx % 2 == 0 and col_idx != 7:
                cell.fill = PatternFill(start_color=COLOR_FILA_PAR_PURP, fill_type="solid")
                
        row_idx += 1
        
    autoajustar_columnas(ws_prod, min_width=14, max_width=45)

    # -------------------------------------------------------------
    # HOJA 4: VENTAS POR COMBO
    # -------------------------------------------------------------
    ws_combos = wb.create_sheet(title="Ventas por Combo")
    ws_combos.views.sheetView[0].showGridLines = True
    
    headers_combos = [
        "Combo", "Precio Combo ($ COP)", "Unidades Vendidas", "Ingresos Generados ($ COP)", "Estado"
    ]
    
    ws_combos.row_dimensions[1].height = 26
    for col_idx, h_text in enumerate(headers_combos, start=1):
        c = ws_combos.cell(row=1, column=col_idx)
        aplicar_estilo_encabezado(c, h_text, bg_color=COLOR_PRIMARIO)
        c.border = thin_border
        
    row_idx = 2
    for item in combos_metricas:
        ws_combos.row_dimensions[row_idx].height = 20
        datos_fila = [
            item.get('nombre', ''),
            item.get('valor_combo', 0.0),
            item.get('unidades_vendidas', 0),
            item.get('ingresos_generados', 0.0),
            "Activo" if item.get('activo', True) else "Inactivo"
        ]
        
        for col_idx, val in enumerate(datos_fila, start=1):
            cell = ws_combos.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            
            if col_idx in (2, 4):
                cell.number_format = "$ #,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx == 3:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.font = Font(name="Segoe UI", size=10, bold=True)
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
            if row_idx % 2 == 0:
                cell.fill = PatternFill(start_color=COLOR_FILA_PAR, fill_type="solid")
                
        row_idx += 1
        
    autoajustar_columnas(ws_combos, min_width=14, max_width=45)

    # -------------------------------------------------------------
    # HOJA 5: ALERTAS DE INVENTARIO Y STOCK BAJO
    # -------------------------------------------------------------
    ws_alertas = wb.create_sheet(title="Alertas de Stock")
    ws_alertas.views.sheetView[0].showGridLines = True
    
    headers_alertas = [
        "Producto", "Categoría", "Stock Disponible", "Valor Unitario ($ COP)", "Nivel de Alerta", "Acción Sugerida"
    ]
    
    ws_alertas.row_dimensions[1].height = 26
    for col_idx, h_text in enumerate(headers_alertas, start=1):
        c = ws_alertas.cell(row=1, column=col_idx)
        aplicar_estilo_encabezado(c, h_text, bg_color="C62828")
        c.border = thin_border

    row_idx = 2
    for prod in alertas_stock:
        ws_alertas.row_dimensions[row_idx].height = 20
        stock_val = prod.stock
        if stock_val == 0:
            alerta_txt = "🚨 AGOTADO"
            accion = "Reabastecer inmediatamente con proveedores"
        elif stock_val <= 3:
            alerta_txt = "⚠️ CRÍTICO (<= 3)"
            accion = "Ordenar reposición de stock urgente"
        else:
            alerta_txt = "⚡ BAJO (<= 5)"
            accion = "Monitorear rotación y programar pedido"

        datos_fila = [
            prod.nombre,
            prod.id_categoria.nombre if prod.id_categoria else "Sin categoría",
            stock_val,
            prod.valor_unitario,
            alerta_txt,
            accion
        ]
        
        for col_idx, val in enumerate(datos_fila, start=1):
            cell = ws_alertas.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            
            if col_idx == 3:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.font = Font(name="Segoe UI", size=10, bold=True, color=COLOR_ALERTA_TXT)
            elif col_idx == 4:
                cell.number_format = "$ #,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(name="Segoe UI", size=10, bold=True, color=COLOR_ALERTA_TXT)
                cell.fill = PatternFill(start_color=COLOR_ALERTA, fill_type="solid")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            if row_idx % 2 == 0 and col_idx != 5:
                cell.fill = PatternFill(start_color="FFF3E0", fill_type="solid")
                
        row_idx += 1
        
    autoajustar_columnas(ws_alertas, min_width=14, max_width=50)

    # Generar buffer en memoria
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"Reporte_Ventas_Sorpresas_Kawaii_{timestamp_str}.xlsx"
    
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
